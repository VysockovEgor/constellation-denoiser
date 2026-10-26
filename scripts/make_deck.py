"""One figure per white 16:9 slide, centred and scaled to fit (contain), no text by default.

    python scripts/make_deck.py [--manifest paper/figures_manifest.csv | --dir <folder>]
                                [--order a.png,b.png] [--out slides/deck_<date>.pptx] [--titles]

Order: manifest rows (fig_no order) when --manifest is used (default), the list
given by --order, otherwise file name. With --titles each slide gets a title
from the manifest column ``title`` or, if there is none, the file name.
Manifest rows whose file does not exist are reported and skipped; figures in
experiments/*/figures that the manifest does not list are reported too.
"""

from __future__ import annotations

import argparse
import csv
import sys
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
SLIDE_W, SLIDE_H = 13.333, 7.5          # inches, 16:9
MARGIN, TITLE_H = 0.25, 0.7


def from_manifest(path: Path) -> list[tuple[Path, str | None]]:
    with path.open(encoding="utf-8") as handle:
        rows = list(csv.DictReader(handle))
    if rows and "fig_no" in rows[0]:
        rows.sort(key=lambda r: float(r["fig_no"]) if r["fig_no"].strip() else float("inf"))
    items, missing = [], []
    for row in rows:
        figure = ROOT / row["file"].strip()
        (items if figure.is_file() else missing).append((figure, (row.get("title") or "").strip() or None))
    for figure, _ in missing:
        print(f"нет файла, пропускаю: {figure.relative_to(ROOT)}", file=sys.stderr)
    listed = {figure.resolve() for figure, _ in items + missing}
    unlisted = sorted(p for p in ROOT.glob("experiments/*/figures/*.png") if p.resolve() not in listed)
    for figure in unlisted:
        print(f"не в манифесте: {figure.relative_to(ROOT)}", file=sys.stderr)
    return items


def build(items: list[tuple[Path, str | None]], out: Path, titles: bool) -> int:
    from PIL import Image
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt

    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(SLIDE_W), Inches(SLIDE_H)
    blank = deck.slide_layouts[6]
    for figure, title in items:
        slide = deck.slides.add_slide(blank)
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)
        top = MARGIN
        if titles:
            box = slide.shapes.add_textbox(Inches(MARGIN), Inches(MARGIN), Inches(SLIDE_W - 2 * MARGIN), Inches(TITLE_H))
            box.text_frame.text = title or figure.name
            box.text_frame.paragraphs[0].runs[0].font.size = Pt(24)
            top += TITLE_H
        with Image.open(figure) as image:
            width, height = image.size
        area_w, area_h = SLIDE_W - 2 * MARGIN, SLIDE_H - top - MARGIN
        scale = min(area_w / width, area_h / height)             # contain: no distortion, no cropping
        pic_w, pic_h = width * scale, height * scale
        left = (SLIDE_W - pic_w) / 2
        pic_top = top + (area_h - pic_h) / 2
        slide.shapes.add_picture(str(figure), Inches(left), Inches(pic_top), Inches(pic_w), Inches(pic_h))
    out.parent.mkdir(parents=True, exist_ok=True)
    deck.save(out)
    return len(items)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    source = parser.add_mutually_exclusive_group()
    source.add_argument("--manifest", type=Path, default=None, help="default: paper/figures_manifest.csv")
    source.add_argument("--dir", type=Path, help="take every *.png of this folder instead of the manifest")
    parser.add_argument("--order", help="comma-separated file names in slide order (overrides the default order)")
    parser.add_argument("--out", type=Path, default=ROOT / "slides" / f"deck_{date.today():%Y%m%d}.pptx")
    parser.add_argument("--titles", action="store_true", help="add a title (manifest 'title' or file name)")
    args = parser.parse_args()

    if args.dir:
        items = [(p, None) for p in sorted(args.dir.glob("*.png"))]
    else:
        manifest = args.manifest or ROOT / "paper" / "figures_manifest.csv"
        items = from_manifest(manifest) if manifest.is_file() else \
            [(p, None) for p in sorted(ROOT.glob("experiments/*/figures/*.png"))]
    if args.order:
        by_name = {figure.name: (figure, title) for figure, title in items}
        wanted = [name.strip() for name in args.order.split(",") if name.strip()]
        missing = [name for name in wanted if name not in by_name]
        if missing:
            parser.error(f"нет среди источников: {', '.join(missing)}")
        items = [by_name[name] for name in wanted]
    if not items:
        parser.error("нет ни одного рисунка")
    count = build(items, args.out, args.titles)
    print(f"{count} слайдов -> {args.out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
